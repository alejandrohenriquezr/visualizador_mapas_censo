# -*- coding: utf-8 -*-
"""Genera la descarga Excel con columnas dinámicas y metadatos auditables."""
from copy import copy
from datetime import datetime
from io import BytesIO
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from dictionary import VARIABLES, es_categoria_invalida
from methodology import universo_formula_explicitos


def _nombre_territorio(nivel, codigo):
    for clave, nombre in VARIABLES.get("geografia", {}).get(nivel, {}).items():
        if str(clave).isdigit() and str(codigo).isdigit() and int(clave) == int(codigo):
            return nombre
    return str(codigo or "")


def _descripcion_variable(tabla, variable):
    info = VARIABLES.get(tabla, {}).get("variables", {}).get(variable, {})
    texto = str(info.get("descripcion") or variable or "").strip()
    texto = re.sub(r"^\s*\d+(?:\.\d+)?[a-z]?\.\s*", "", texto, flags=re.I)
    return texto.strip(" ¿?\t\n")


def _codigo_region(nivel, codigo):
    if codigo in (None, ""):
        return None
    try:
        entero = int(codigo)
    except (TypeError, ValueError):
        return None
    if nivel == "region":
        return entero
    if nivel == "comuna":
        return int(str(entero).zfill(5)[:2])
    if nivel == "provincia":
        return int(str(entero).zfill(3)[:2])
    return None


def _region_primaria(respuesta, propiedades):
    interp = respuesta.get("interpretacion") or {}
    filtro = interp.get("filtro_geografico_nivel")
    codigo_filtro = interp.get("filtro_geografico_codigo")
    nivel = respuesta.get("nivel_geografico") or interp.get("nivel_geografico") or "region"

    if filtro == "region" and codigo_filtro is not None:
        codigo_region = int(codigo_filtro)
    elif filtro in {"comuna", "provincia"} and codigo_filtro is not None:
        codigo_region = _codigo_region(filtro, codigo_filtro)
    else:
        codigo_region = _codigo_region(nivel, propiedades.get("codigo"))

    return "Región", _nombre_territorio("region", codigo_region) if codigo_region is not None else ""


def _usa_secundario(respuesta):
    interp = respuesta.get("interpretacion") or {}
    nivel = respuesta.get("nivel_geografico") or interp.get("nivel_geografico")
    return nivel == "comuna" or interp.get("filtro_geografico_nivel") == "comuna"


def _comuna_secundaria(respuesta, propiedades):
    interp = respuesta.get("interpretacion") or {}
    filtro = interp.get("filtro_geografico_nivel")
    codigo_filtro = interp.get("filtro_geografico_codigo")
    nivel = respuesta.get("nivel_geografico") or interp.get("nivel_geografico")
    if filtro == "comuna" and codigo_filtro is not None:
        return "Comuna", _nombre_territorio("comuna", codigo_filtro)
    if nivel == "comuna":
        return "Comuna", propiedades.get("nombre") or _nombre_territorio("comuna", propiedades.get("codigo"))
    return None, None


def _etiqueta_categoria(tabla, variable, valores, etiqueta=None):
    if etiqueta:
        return str(etiqueta)
    info = VARIABLES.get(tabla, {}).get("variables", {}).get(variable, {})
    cats = info.get("categorias", {})
    return " / ".join(str(cats.get(str(v), v)) for v in (valores or []))


def _hojas_ast(nodo):
    if not isinstance(nodo, dict):
        return []
    if nodo.get("op") in {"and", "or"}:
        salida = []
        for item in nodo.get("args") or []:
            salida.extend(_hojas_ast(item))
        return salida
    if nodo.get("op") == "not":
        hojas = _hojas_ast(nodo.get("arg"))
        for item in hojas:
            item = dict(item); item["_negado"] = True
        return hojas
    return [nodo]


def _categorias_fijas(respuesta):
    """Categorías de filtros fijos que deben acompañar cada fila exportada."""
    interp = respuesta.get("interpretacion") or {}
    plan = interp.get("plan_cruce") or {}
    salida = []
    filtros_fuente = (_hojas_ast(plan.get("filtro_ast"))
                      if plan.get("filtro_ast") else (plan.get("filtros") or []))
    for filtro in filtros_fuente:
        if filtro.get("tipo") != "categoria":
            continue
        tabla, variable = filtro.get("tabla"), filtro.get("variable")
        etiqueta = _etiqueta_categoria(tabla, variable, filtro.get("valores"), filtro.get("etiqueta"))
        if etiqueta:
            prefijo = "Excluye " if filtro.get("_negado") or filtro.get("modo") == "not_in" else ""
            salida.append(f"{prefijo}{_descripcion_variable(tabla, variable)}: {etiqueta}")

    if not plan and interp.get("categoria_valor") is not None:
        tabla, variable = interp.get("tabla"), interp.get("variable")
        valores = interp.get("categoria_valores") or [interp.get("categoria_valor")]
        etiqueta = _etiqueta_categoria(tabla, variable, valores)
        if etiqueta:
            salida.append(f"{_descripcion_variable(tabla, variable)}: {etiqueta}")
    return salida


def _categoria_compuesta(respuesta, item=None, tipo=None):
    partes = list(_categorias_fijas(respuesta))
    item = item or {}
    etiqueta_item = item.get("etiqueta")
    if tipo == "barras_mapa" and etiqueta_item:
        sev = respuesta.get("severidad_etiqueta")
        if sev:
            partes.append(f"{etiqueta_item}: {sev}")
        else:
            partes.append(str(etiqueta_item))
    elif etiqueta_item:
        # Los cruces ya entregan etiquetas del tipo «Sexo: Hombre · ...».
        partes.append(str(etiqueta_item))

    unicas = []
    for parte in partes:
        if parte and parte not in unicas:
            unicas.append(parte)
    return " | ".join(unicas) if unicas else None


def _registros_resultado(respuesta):
    """Devuelve registros normalizados y las columnas de resultados disponibles."""
    tipo = respuesta.get("tipo_visualizacion")
    tortas = tipo == "tortas_mapa"
    barras = tipo == "barras_mapa"
    registros = []
    columnas_resultado = []

    def geo_base(propiedades):
        nivel_p, nombre_p = _region_primaria(respuesta, propiedades)
        base = {
            "Nivel geográfico primario": nivel_p,
            "Nombre del territorio primario": nombre_p,
        }
        if _usa_secundario(respuesta):
            nivel_s, nombre_s = _comuna_secundaria(respuesta, propiedades)
            base["Nivel geográfico secundario (opcional)"] = nivel_s
            base["Nombre del territorio secundario"] = nombre_s
        return base

    if barras:
        metrica = respuesta.get("metrica")
        if metrica == "porcentaje":
            columnas_resultado = ["Cantidad", "Porcentaje", "Total"]
        elif metrica in {"promedio", "mediana", "suma", "minimo", "maximo"}:
            columnas_resultado = ["Valor del indicador/tasa/porcentaje"]
        else:
            columnas_resultado = ["Cantidad", "Porcentaje", "Total"]
        for territorio in respuesta.get("datos") or []:
            base = geo_base(territorio)
            for item in territorio.get("barras") or []:
                fila = dict(base)
                fila["Categoría"] = _categoria_compuesta(respuesta, item, "barras_mapa")
                if metrica in {"promedio", "mediana", "suma", "minimo", "maximo"}:
                    fila["Valor del indicador/tasa/porcentaje"] = item.get("valor")
                else:
                    fila["Cantidad"] = item.get("conteo")
                    fila["Porcentaje"] = item.get("porcentaje")
                    fila["Total"] = item.get("denominador")
                registros.append(fila)
        return registros, columnas_resultado

    datos = respuesta.get("datos")
    if datos is None:
        datos = [feature.get("properties") or {}
                 for feature in (respuesta.get("geojson") or {}).get("features", [])]

    if tortas:
        columnas_resultado = ["Cantidad", "Porcentaje", "Total"]
        for territorio in datos or []:
            base = geo_base(territorio)
            total = territorio.get("total")
            # V33: el mapa puede agrupar en “Otros”, pero Excel conserva todas
            # las categorías válidas entregadas por el motor.
            distribucion_excel = territorio.get("distribucion_completa") or territorio.get("distribucion") or []
            for item in distribucion_excel:
                fila = dict(base)
                fila["Categoría"] = _categoria_compuesta(respuesta, item, "tortas_mapa")
                fila["Cantidad"] = item.get("valor")
                fila["Porcentaje"] = item.get("porcentaje")
                fila["Total"] = total
                registros.append(fila)
        return registros, columnas_resultado

    interp = respuesta.get("interpretacion") or {}
    operacion = interp.get("operacion") or "conteo"
    if operacion == "conteo":
        col_valor = "Cantidad"
    elif operacion in {"porcentaje", "porcentaje_rango"}:
        col_valor = "Porcentaje"
    else:
        col_valor = "Valor del indicador/tasa/porcentaje"
    columnas_resultado = [col_valor]
    categoria = _categoria_compuesta(respuesta)
    for territorio in datos or []:
        fila = geo_base(territorio)
        if categoria:
            fila["Categoría"] = categoria
        fila[col_valor] = territorio.get("valor")
        registros.append(fila)
    return registros, columnas_resultado


def _variables_involucradas(interp):
    items = []
    plan = interp.get("plan_cruce") or {}
    if plan:
        for d in plan.get("dimensiones") or []:
            items.append((d.get("tabla"), d.get("variable")))
        filtros_fuente = (_hojas_ast(plan.get("filtro_ast"))
                          if plan.get("filtro_ast") else (plan.get("filtros") or []))
        for f in filtros_fuente:
            if f.get("tabla") and f.get("variable"):
                items.append((f.get("tabla"), f.get("variable")))
            for sf in f.get("subfiltros") or []:
                if sf.get("tabla") and sf.get("variable"):
                    items.append((sf.get("tabla"), sf.get("variable")))
        medida = plan.get("medida") or {}
        if medida.get("tabla") and medida.get("variable"):
            items.append((medida.get("tabla"), medida.get("variable")))
    else:
        items.append((interp.get("tabla"), interp.get("variable")))
        for variable in interp.get("variables_dimensiones") or []:
            items.append(("personas", variable))
    unicos = []
    for item in items:
        if item[0] and item[1] and item not in unicos:
            unicos.append(item)
    return unicos


def _exclusiones_calculo(interp):
    lineas = []
    for tabla, variable in _variables_involucradas(interp):
        info = VARIABLES.get(tabla, {}).get("variables", {}).get(variable, {})
        excluidas = [
            f"{codigo}: {etiqueta}"
            for codigo, etiqueta in info.get("categorias", {}).items()
            if es_categoria_invalida(codigo, etiqueta)
        ]
        if excluidas:
            lineas.append(
                f"{_descripcion_variable(tabla, variable)} — categorías/códigos excluidos: "
                + "; ".join(excluidas)
            )
    return lineas


def _copiar_estilo_fila(ws, origen, destino, columnas=8):
    for columna in range(1, columnas + 1):
        fuente, celda = ws.cell(origen, columna), ws.cell(destino, columna)
        celda._style = copy(fuente._style)
        celda.font = copy(fuente.font)
        celda.fill = copy(fuente.fill)
        celda.border = copy(fuente.border)
        celda.alignment = copy(fuente.alignment)
        celda.number_format = fuente.number_format


def nombre_archivo(titulo):
    texto = unicodedata.normalize("NFKD", titulo or "resultados_censo")
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"[^A-Za-z0-9]+", "_", texto).strip("_").lower()
    return (texto[:60].rstrip("_") or "resultados_censo") + ".xlsx"


FUENTE_FALLBACK = (
    "Fuente: Censo de Población y Vivienda 2024 - Instituto Nacional de Estadísticas."
)
NOTA_FALLBACK = (
    "Nota: través de una consulta realizada por el usuario en lenguaje natural e "
    "interpretada por un LLM, por lo que los resultados deben ser analizados con "
    "detención y deben ser contrastados con información oficial publicada en "
    "www.ine.gob.cl o generar nuevos resultados en https://redatam-ine.ine.cl."
)


def _buscar_fila_prefijo(ws, prefijo, excluir=None):
    prefijo = prefijo.lower()
    excluir = tuple((x or "").lower() for x in (excluir or []))
    for fila in range(1, ws.max_row + 1):
        valor = ws.cell(fila, 1).value
        if not isinstance(valor, str):
            continue
        limpio = valor.strip().lower()
        if limpio.startswith(prefijo) and not any(limpio.startswith(x) for x in excluir):
            return fila, valor
    return None, None


def _capturar_estilo_celda(celda):
    return {
        "_style": copy(celda._style),
        "font": copy(celda.font),
        "fill": copy(celda.fill),
        "border": copy(celda.border),
        "alignment": copy(celda.alignment),
        "number_format": celda.number_format,
        "protection": copy(celda.protection),
    }


def _aplicar_estilo_celda(celda, estilo, alineacion=None):
    if estilo:
        celda._style = copy(estilo["_style"])
        celda.font = copy(estilo["font"])
        celda.fill = copy(estilo["fill"])
        celda.border = copy(estilo["border"])
        celda.alignment = copy(estilo["alignment"])
        celda.number_format = estilo["number_format"]
        celda.protection = copy(estilo["protection"])
    if alineacion:
        ali = copy(celda.alignment)
        ali.horizontal = alineacion
        ali.vertical = ali.vertical or "top"
        ali.wrap_text = True
        celda.alignment = ali


def _capturar_pie_resultado(ws):
    fila_fuente, texto_fuente = _buscar_fila_prefijo(ws, "fuente:")
    fila_nota, texto_nota = _buscar_fila_prefijo(ws, "nota:", excluir=("nota final",))
    fila_final, texto_final = _buscar_nota_final(ws)

    def estilo(fila, respaldo):
        return _capturar_estilo_celda(ws.cell(fila, 1)) if fila else respaldo

    # Si la plantilla no contiene Fuente/Nota, se usa el estilo de Nota final
    # como respaldo, sin inventar un formato distinto.
    estilo_respaldo = _capturar_estilo_celda(ws.cell(fila_final, 1)) if fila_final else None
    return {
        "fuente": texto_fuente or FUENTE_FALLBACK,
        "nota": texto_nota or NOTA_FALLBACK,
        "nota_final": texto_final or "Nota final",
        "estilo_fuente": estilo(fila_fuente, estilo_respaldo),
        "estilo_nota": estilo(fila_nota, estilo_respaldo),
        "estilo_final": estilo(fila_final, estilo_respaldo),
        "altura_fuente": ws.row_dimensions[fila_fuente].height if fila_fuente else None,
        "altura_nota": ws.row_dimensions[fila_nota].height if fila_nota else None,
    }


def _alinear_dato(celda, valor):
    ali = copy(celda.alignment)
    es_numero = isinstance(valor, (int, float)) and not isinstance(valor, bool)
    ali.horizontal = "right" if es_numero else "left"
    ali.vertical = ali.vertical or "center"
    celda.alignment = ali


def _buscar_nota_final(ws):
    for fila in range(1, min(ws.max_row, 80) + 1):
        valor = ws.cell(fila, 1).value
        if isinstance(valor, str) and "nota final" in valor.lower():
            return fila, valor
    return 22, "Nota final"





def _universo_formula_exportacion(respuesta):
    """Metadatos metodológicos legibles para la hoja Nota."""
    interp = respuesta.get("interpretacion") or {}
    universo = str(interp.get("universo") or "").strip()
    formula = str(interp.get("formula") or "").strip()
    operacion = str(interp.get("operacion") or "").strip().lower()
    tabla = str(interp.get("entidad_objetivo") or interp.get("tabla") or "registros")
    variable = interp.get("variable")
    descripcion = _descripcion_variable(interp.get("tabla"), variable) if variable else ""

    # Para consultas estadísticas ordinarias, sustituye los textos genéricos
    # por una descripción trazable del universo efectivo y de la operación.
    # Los indicadores que ya traen universo/fórmula explícitos conservan esos
    # metadatos; los productos OD mantienen su tratamiento específico abajo.
    if respuesta.get("tipo_visualizacion") != "matriz_od" and (not universo or not formula):
        universo_exp, formula_exp = universo_formula_explicitos(respuesta)
        if not universo and universo_exp:
            universo = universo_exp
        if not formula and formula_exp:
            formula = formula_exp

    if not universo:
        if respuesta.get("tipo_visualizacion") == "matriz_od":
            dominio = (respuesta.get("config_od") or {}).get("dominio")
            if dominio == "migracion_interna":
                universo = ("Personas que ya habían nacido en abril de 2019, residían en Chile en 2019 y "
                            "tienen territorio de origen 2019 identificable.")
            else:
                universo = "Personas ocupadas con lugar de trabajo territorialmente identificable en Chile."
        elif operacion in {"porcentaje", "porcentaje_rango", "distribucion", "razon"} and descripcion:
            universo = f"{tabla.capitalize()} con información válida para «{descripcion}» y que cumplen los filtros de la consulta."
        else:
            universo = f"{tabla.capitalize()} que cumplen los filtros definidos por la consulta."

    if not formula:
        if operacion in {"porcentaje", "porcentaje_rango"}:
            formula = "Porcentaje = (casos que cumplen la condición / casos válidos del universo) × 100"
        elif operacion == "distribucion":
            formula = "Porcentaje de cada categoría = (casos de la categoría / casos válidos del universo) × 100"
        elif operacion == "promedio":
            formula = "Promedio = suma de los valores válidos / número de casos con valor válido"
        elif operacion == "razon":
            factor = interp.get("factor")
            factor_txt = f" × {factor}" if factor not in (None, 1, 1.0) else ""
            formula = f"Indicador = numerador / denominador{factor_txt}"
        elif respuesta.get("tipo_visualizacion") == "matriz_od":
            metrica = (respuesta.get("config_od") or {}).get("metrica") or "cantidad"
            if metrica == "porcentaje_origen":
                formula = "Porcentaje por origen = (flujo origen→destino / total de flujos del origen) × 100"
            elif metrica == "porcentaje_destino":
                formula = "Porcentaje por destino = (flujo origen→destino / total de flujos del destino) × 100"
            else:
                formula = "Cantidad = conteo de personas que cumplen el par origen→destino seleccionado"
        elif operacion == "conteo":
            formula = "Conteo directo de casos que cumplen los filtros de la consulta"
        else:
            formula = "Cálculo directo según la operación seleccionada"
    return universo, formula


def _escribir_universo_formula(nota, respuesta):
    """Escribe Universo y Fórmula en las celdas metodológicas solicitadas."""
    universo, formula = _universo_formula_exportacion(respuesta)
    nota["A20"] = None
    nota["A21"] = None
    nota["B20"] = "Universo"
    nota["E20"] = universo
    nota["B21"] = "Fórmula"
    nota["E21"] = formula
    nota["B20"].font = Font(bold=True)
    nota["E21"].font = Font(bold=True)
    nota["E20"].alignment = Alignment(wrap_text=True, vertical="top")
    nota["E21"].alignment = Alignment(wrap_text=True, vertical="top")

def _crear_excel_od(plantilla, respuesta, pregunta):
    """Exporta cualquier producto origen-destino en formato matricial y largo."""
    if not plantilla.exists():
        raise FileNotFoundError(f"No se encontró la plantilla Excel en '{plantilla}'.")
    wb = load_workbook(plantilla)
    if "Nota" not in wb.sheetnames:
        raise ValueError("La plantilla debe contener la hoja Nota.")

    # La hoja Resultado genérica se reemplaza por las dos estructuras OD.
    if "Resultado" in wb.sheetnames:
        wb.remove(wb["Resultado"])
    for nombre in ("Matriz_OD", "Flujos_OD"):
        if nombre in wb.sheetnames:
            wb.remove(wb[nombre])
    matriz = wb.create_sheet("Matriz_OD")
    flujos_ws = wb.create_sheet("Flujos_OD")
    nota = wb["Nota"]

    interp = respuesta.get("interpretacion") or {}
    config = respuesta.get("config_od") or {}
    dominio = config.get("dominio") or "movilidad_laboral"
    unidad = config.get("unidad") or "comuna"
    unidad_cap = unidad.capitalize()
    origen_etiqueta = config.get("origen_etiqueta") or f"{unidad_cap} de residencia"
    destino_etiqueta = config.get("destino_etiqueta") or f"{unidad_cap} de trabajo"
    titulo = interp.get("descripcion") or ("Matriz de migración interna 2019–2024" if dominio == "migracion_interna" else "Matriz de movilidad laboral entre comunas")
    nota["B5"] = "Fecha de obtención del cuadro: " + datetime.now().strftime("%d-%m-%Y %H:%M")
    for fila in range(14, max(nota.max_row, 24) + 1):
        nota.cell(fila, 2).value = None
        nota.cell(fila, 5).value = None
    if dominio == "migracion_interna":
        producto = f"Matriz origen-destino de migración interna 2019–2024 por {unidad}"
        metrica_meta = interp.get("migracion_metrica") or config.get("metrica") or "cantidad"
        diagonal_meta = bool(interp.get("migracion_incluir_diagonal", config.get("incluir_diagonal", False)))
        universo = "Personas que ya habían nacido en abril de 2019, residían en Chile y tienen territorio de origen 2019 identificable"
    else:
        producto = "Matriz origen-destino comunal de movilidad laboral"
        metrica_meta = interp.get("movilidad_metrica") or config.get("metrica") or "cantidad"
        diagonal_meta = bool(interp.get("movilidad_incluir_diagonal", config.get("incluir_diagonal", True)))
        universo = "Personas ocupadas con un único lugar de trabajo comunal identificable en Chile"
    metadatos = [
        ("Pregunta", pregunta),
        ("Producto", producto),
        ("Métrica mostrada", metrica_meta),
        (f"Incluye misma {unidad}", "Sí" if diagonal_meta else "No"),
        ("Universo", universo),
        ("Entidad objetivo", interp.get("entidad_objetivo") or interp.get("tabla") or "personas"),
    ]
    for fila, (nombre, valor) in enumerate(metadatos, start=14):
        nota.cell(fila, 2, nombre); nota.cell(fila, 5, valor)
    _escribir_universo_formula(nota, respuesta)

    flujos = list(respuesta.get("flujos_od") or [])
    metrica = config.get("metrica") or interp.get("movilidad_metrica") or interp.get("migracion_metrica") or "cantidad"
    clave_valor = {
        "cantidad": "cantidad",
        "porcentaje_origen": "porcentaje_origen",
        "porcentaje_destino": "porcentaje_destino",
    }.get(metrica, "cantidad")
    etiqueta_valor = {
        "cantidad": "Cantidad de personas" if dominio == "migracion_interna" else "Cantidad de personas ocupadas",
        "porcentaje_origen": f"% sobre {origen_etiqueta.lower()}",
        "porcentaje_destino": f"% sobre {destino_etiqueta.lower()}",
    }.get(metrica, "Cantidad")

    azul = "174A7E"; celeste = "D9EAF7"; borde = Side(style="thin", color="D0D7DE")
    header_fill = PatternFill("solid", fgColor=azul); light_fill = PatternFill("solid", fgColor=celeste)
    header_font = Font(bold=True, color="FFFFFF"); bold = Font(bold=True)

    # -------------------- Hoja larga --------------------
    flujos_ws["A1"] = titulo
    flujos_ws.merge_cells("A1:J1")
    flujos_ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    flujos_ws["A1"].fill = header_fill
    headers = [
        "Región origen", origen_etiqueta, "Código origen",
        "Región destino", destino_etiqueta, "Código destino",
        "Cantidad", "% origen", "% destino", f"Misma {unidad}",
    ]
    for c, h in enumerate(headers, 1):
        cell=flujos_ws.cell(3,c,h); cell.fill=header_fill; cell.font=header_font; cell.alignment=Alignment(horizontal="center", wrap_text=True); cell.border=Border(bottom=borde)
    for r, f in enumerate(flujos, 4):
        vals=[f.get("origen_region_nombre"), f.get("origen_nombre"), f.get("origen_codigo"),
              f.get("destino_region_nombre"), f.get("destino_nombre"), f.get("destino_codigo"),
              f.get("cantidad"), f.get("porcentaje_origen"), f.get("porcentaje_destino"),
              "Sí" if f.get("misma_comuna") else "No"]
        for c,v in enumerate(vals,1):
            cell=flujos_ws.cell(r,c,v); cell.border=Border(bottom=Border().bottom); cell.alignment=Alignment(horizontal="right" if c in (3,6,7,8,9) else "left")
            if r%2==0: cell.fill=light_fill
        flujos_ws.cell(r,8).number_format='0.00%'.replace('%','')+'"%"'
        flujos_ws.cell(r,9).number_format='0.00%'.replace('%','')+'"%"'
    flujos_ws.freeze_panes="A4"
    flujos_ws.auto_filter.ref=f"A3:J{max(3,3+len(flujos))}"
    widths=[24,28,22,24,28,20,14,14,14,14]
    for i,w in enumerate(widths,1): flujos_ws.column_dimensions[get_column_letter(i)].width=w

    # Fuente/Nota inmediatamente después de los flujos, igual que el exportador general.
    pie_largo = 4 + len(flujos)
    for fila, texto in ((pie_largo, FUENTE_FALLBACK), (pie_largo + 1, NOTA_FALLBACK),
                        (pie_largo + 2, "Nota final: " + str(interp.get("nota") or respuesta.get("nota") or ""))):
        flujos_ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=10)
        flujos_ws.cell(fila, 1, texto); flujos_ws.cell(fila, 1).alignment=Alignment(wrap_text=True, horizontal="left")
        if fila == pie_largo: flujos_ws.cell(fila, 1).font=bold

    # -------------------- Hoja matriz --------------------
    matriz["A1"] = titulo
    matriz.merge_cells("A1:F1")
    matriz["A1"].font=Font(bold=True,size=14,color="FFFFFF"); matriz["A1"].fill=header_fill
    matriz["A2"] = f"Celdas: {etiqueta_valor}. Filas = {origen_etiqueta.lower()}; columnas = {destino_etiqueta.lower()}."
    matriz.merge_cells("A2:F2")
    matriz["A2"].alignment=Alignment(wrap_text=True)

    # Orden por total descendente para que la matriz sea útil desde la esquina superior izquierda.
    ori_tot={}; des_tot={}
    for f in flujos:
        ori_tot[f["origen_codigo"]]=ori_tot.get(f["origen_codigo"],0)+int(f.get("cantidad") or 0)
        des_tot[f["destino_codigo"]]=des_tot.get(f["destino_codigo"],0)+int(f.get("cantidad") or 0)
    nombres_o={f["origen_codigo"]:f["origen_nombre"] for f in flujos}
    nombres_d={f["destino_codigo"]:f["destino_nombre"] for f in flujos}
    origenes=sorted(ori_tot, key=lambda c:(-ori_tot[c], nombres_o.get(c,"")))
    destinos=sorted(des_tot, key=lambda c:(-des_tot[c], nombres_d.get(c,"")))
    lookup={(f["origen_codigo"],f["destino_codigo"]):f for f in flujos}

    start_row=4; start_col=2
    matriz.cell(start_row,1,f"{origen_etiqueta} ↓ / {destino_etiqueta} →")
    matriz.cell(start_row,1).fill=header_fill; matriz.cell(start_row,1).font=header_font; matriz.cell(start_row,1).alignment=Alignment(wrap_text=True)
    for j,cod in enumerate(destinos,start_col):
        cell=matriz.cell(start_row,j,nombres_d.get(cod,str(cod))); cell.fill=header_fill; cell.font=header_font; cell.alignment=Alignment(text_rotation=90,horizontal="center",vertical="bottom")
    total_col=start_col+len(destinos)
    matriz.cell(start_row,total_col,"Total cantidad origen").fill=header_fill; matriz.cell(start_row,total_col).font=header_font

    for i,cod_o in enumerate(origenes,start_row+1):
        h=matriz.cell(i,1,nombres_o.get(cod_o,str(cod_o))); h.fill=header_fill; h.font=header_font
        for j,cod_d in enumerate(destinos,start_col):
            f=lookup.get((cod_o,cod_d)); val=None if not f else f.get(clave_valor)
            cell=matriz.cell(i,j,val)
            if metrica.startswith("porcentaje") and val is not None: cell.number_format='0.00"%"'
            elif val is not None: cell.number_format='#,##0'
            if cod_o==cod_d and f is not None: cell.fill=light_fill
            cell.alignment=Alignment(horizontal="right")
        matriz.cell(i,total_col,ori_tot.get(cod_o,0)).number_format='#,##0'
        matriz.cell(i,total_col).font=bold

    total_row=start_row+1+len(origenes)
    matriz.cell(total_row,1,"Total cantidad destino").fill=header_fill; matriz.cell(total_row,1).font=header_font
    for j,cod_d in enumerate(destinos,start_col):
        matriz.cell(total_row,j,des_tot.get(cod_d,0)).number_format='#,##0'; matriz.cell(total_row,j).font=bold
    matriz.cell(total_row,total_col,sum(ori_tot.values())).number_format='#,##0'; matriz.cell(total_row,total_col).font=bold
    matriz.freeze_panes="B5"
    matriz.column_dimensions['A'].width=34
    for j in range(start_col,total_col): matriz.column_dimensions[get_column_letter(j)].width=13
    matriz.column_dimensions[get_column_letter(total_col)].width=20
    matriz.row_dimensions[start_row].height=110

    # Fuente y nota metodológica debajo de la matriz.
    pie=total_row+2
    matriz.cell(pie,1,FUENTE_FALLBACK); matriz.cell(pie,1).font=bold
    matriz.merge_cells(start_row=pie,start_column=1,end_row=pie,end_column=max(2,total_col))
    matriz.cell(pie+1,1,NOTA_FALLBACK); matriz.merge_cells(start_row=pie+1,start_column=1,end_row=pie+1,end_column=max(2,total_col)); matriz.cell(pie+1,1).alignment=Alignment(wrap_text=True)
    matriz.cell(pie+2,1,interp.get("nota") or respuesta.get("nota") or ""); matriz.merge_cells(start_row=pie+2,start_column=1,end_row=pie+2,end_column=max(2,total_col)); matriz.cell(pie+2,1).alignment=Alignment(wrap_text=True)

    salida=BytesIO(); wb.save(salida); salida.seek(0)
    return salida, nombre_archivo(titulo)


def _orden_edad_piramide(config, codigo):
    orden = [str(x) for x in config.get("orden_edad") or []]
    codigo = str(codigo)
    if codigo in orden:
        return (0, orden.index(codigo))
    try:
        return (1, float(codigo))
    except (TypeError, ValueError):
        m = re.match(r"\s*(\d+)", codigo)
        return (2, int(m.group(1)) if m else 10**9, codigo)


def _filas_piramides(respuesta):
    config = respuesta.get("piramide") or {}
    if not config.get("disponible"):
        return []
    try:
        i_edad = int(config["indice_edad"])
        i_sexo = int(config["indice_sexo"])
    except (KeyError, TypeError, ValueError):
        return []
    hombre = str(config.get("sexo_hombre_codigo", "1"))
    mujer = str(config.get("sexo_mujer_codigo", "2"))
    etiquetas = {str(k): str(v) for k, v in (config.get("etiquetas_edad") or {}).items()}
    nivel = respuesta.get("nivel_geografico") or config.get("nivel_geografico") or "region"
    filas = []
    for territorio in respuesta.get("datos") or []:
        por_edad = {}
        for item in (territorio.get("distribucion_completa") or territorio.get("distribucion") or []):
            partes = str(item.get("codigo") or "").split("|")
            if len(partes) <= max(i_edad, i_sexo):
                continue
            edad, sexo = partes[i_edad], partes[i_sexo]
            if sexo not in {hombre, mujer}:
                continue
            rec = por_edad.setdefault(edad, {"hombres": 0, "mujeres": 0})
            valor = int(item.get("valor") or 0)
            if sexo == hombre:
                rec["hombres"] += valor
            else:
                rec["mujeres"] += valor
        codigo_territorio = territorio.get("codigo")
        codigo_region = _codigo_region(nivel, codigo_territorio)
        nombre_region = _nombre_territorio("region", codigo_region) if codigo_region is not None else ""
        for edad in sorted(por_edad, key=lambda x: _orden_edad_piramide(config, x)):
            rec = por_edad[edad]
            etiqueta = etiquetas.get(str(edad))
            if not etiqueta:
                try:
                    etiqueta = f"{int(float(edad))} años"
                except (TypeError, ValueError):
                    etiqueta = str(edad)
            filas.append({
                "Nivel geográfico": nivel.capitalize(),
                "Región": nombre_region,
                "Código territorio": codigo_territorio,
                "Territorio": territorio.get("nombre") or str(codigo_territorio or ""),
                "Edad / grupo de edad": etiqueta,
                "Hombres": rec["hombres"],
                "Mujeres": rec["mujeres"],
                "Total": rec["hombres"] + rec["mujeres"],
            })
    return filas


def _agregar_hoja_piramides(wb, respuesta):
    filas = _filas_piramides(respuesta)
    if not filas:
        return
    nombre = "Piramides_edad"
    if nombre in wb.sheetnames:
        wb.remove(wb[nombre])
    ws = wb.create_sheet(nombre)
    headers = ["Nivel geográfico", "Región", "Código territorio", "Territorio",
               "Edad / grupo de edad", "Hombres", "Mujeres", "Total"]
    azul = "174A7E"
    fill = PatternFill("solid", fgColor=azul)
    font = Font(bold=True, color="FFFFFF")
    for c, h in enumerate(headers, 1):
        celda = ws.cell(1, c, h)
        celda.fill = fill; celda.font = font
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, fila in enumerate(filas, 2):
        for c, h in enumerate(headers, 1):
            valor = fila.get(h)
            celda = ws.cell(r, c, valor)
            celda.alignment = Alignment(horizontal="right" if h in {"Código territorio", "Hombres", "Mujeres", "Total"} else "left")
            if h in {"Hombres", "Mujeres", "Total"}:
                celda.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{1 + len(filas)}"
    for col, width in {"A":18, "B":28, "C":18, "D":32, "E":22, "F":16, "G":16, "H":16}.items():
        ws.column_dimensions[col].width = width

def crear_excel(plantilla, respuesta, pregunta, sql):
    if respuesta.get("tipo_visualizacion") == "matriz_od":
        return _crear_excel_od(plantilla, respuesta, pregunta)
    if not plantilla.exists():
        raise FileNotFoundError(f"No se encontró la plantilla Excel en '{plantilla}'.")
    wb = load_workbook(plantilla)
    if not {"Nota", "Resultado"}.issubset(wb.sheetnames):
        raise ValueError("La plantilla debe contener las hojas Nota y Resultado.")

    nota, resultado = wb["Nota"], wb["Resultado"]
    interp = respuesta.get("interpretacion") or {}
    titulo = interp.get("descripcion") or "Resultados de la consulta"
    resultado["A1"] = titulo
    nota["B5"] = "Fecha de obtención del cuadro: " + datetime.now().strftime("%d-%m-%Y %H:%M")

    # La hoja Nota conserva solo metadatos legibles. El plan interno de cruce
    # y la instrucción SQL son detalles de implementación y no se exportan.
    for fila in range(14, max(nota.max_row, 24) + 1):
        nota.cell(fila, 2).value = None
        nota.cell(fila, 5).value = None
    metadatos = [
        ("Pregunta", pregunta),
        ("Tabla", interp.get("tabla")),
        ("Variable", interp.get("variable")),
        ("Operación", interp.get("operacion")),
        ("Nivel geográfico", interp.get("nivel_geografico")),
        ("Entidad objetivo", interp.get("entidad_objetivo") or interp.get("tabla")),
    ]
    for fila, (nombre, valor) in enumerate(metadatos, start=14):
        nota.cell(fila, 2, nombre)
        nota.cell(fila, 5, "" if valor is None else str(valor))
    _escribir_universo_formula(nota, respuesta)

    registros, columnas_resultado = _registros_resultado(respuesta)
    if not registros:
        raise ValueError("No hay resultados para exportar.")

    usa_secundario = _usa_secundario(respuesta)
    usa_categoria = any(r.get("Categoría") not in (None, "") for r in registros)
    encabezados = ["Nivel geográfico primario", "Nombre del territorio primario"]
    if usa_secundario:
        encabezados += ["Nivel geográfico secundario (opcional)", "Nombre del territorio secundario"]
    if usa_categoria:
        encabezados.append("Categoría")
    # Elimina columnas de resultado completamente vacías, pero conserva al menos una.
    for col in columnas_resultado:
        if any(r.get(col) is not None for r in registros):
            encabezados.append(col)
    if not any(h in encabezados for h in columnas_resultado):
        encabezados.append(columnas_resultado[0])

    # Conserva textos y estilos del pie original de la plantilla antes de
    # limpiar/mover filas. Fuente y Nota deben sobrevivir sin cambios.
    pie = _capturar_pie_resultado(resultado)

    # Deshace merges de la zona tabular/pie antes de limpiar; de lo contrario
    # openpyxl devuelve MergedCell de solo lectura.
    for rango in list(resultado.merged_cells.ranges):
        if rango.min_row >= 3:
            resultado.unmerge_cells(str(rango))

    # Limpia la zona tabular de la plantilla, conservando estilos base.
    for fila in range(3, max(resultado.max_row, 30) + 1):
        for columna in range(1, 9):
            resultado.cell(fila, columna).value = None

    # Encabezados y filas dinámicas. Se conserva el patrón visual de la
    # plantilla: fila celeste, fila sin relleno, celeste, sin relleno, ...
    # usando las filas 4 y 5 como patrones alternados.
    for col, encabezado in enumerate(encabezados, start=1):
        celda = resultado.cell(3, col, encabezado)
        _alinear_dato(celda, encabezado)
    for posicion, (indice, registro) in enumerate(enumerate(registros, start=4)):
        fila_patron = 4 if posicion % 2 == 0 else 5
        _copiar_estilo_fila(resultado, fila_patron, indice,
                            columnas=max(8, len(encabezados)))
        for col, encabezado in enumerate(encabezados, start=1):
            valor = registro.get(encabezado)
            celda = resultado.cell(indice, col, valor)
            _alinear_dato(celda, valor)

    # Oculta las columnas opcionales que no fueron necesarias; las columnas
    # utilizadas quedan contiguas y visibles desde A.
    max_plantilla = max(8, resultado.max_column)
    for col in range(1, max_plantilla + 1):
        letra = get_column_letter(col)
        resultado.column_dimensions[letra].hidden = col > len(encabezados)
    anchos = {
        "Nivel geográfico primario": 22,
        "Nombre del territorio primario": 30,
        "Nivel geográfico secundario (opcional)": 28,
        "Nombre del territorio secundario": 30,
        "Categoría": 56,
        "Cantidad": 18,
        "Porcentaje": 18,
        "Total": 18,
        "Valor del indicador/tasa/porcentaje": 30,
    }
    for col, encabezado in enumerate(encabezados, start=1):
        resultado.column_dimensions[get_column_letter(col)].width = anchos.get(encabezado, 20)

    ultima_fila_datos = 3 + len(registros)
    resultado.freeze_panes = "A4"
    resultado.auto_filter.ref = f"A3:{get_column_letter(len(encabezados))}{ultima_fila_datos}"

    # Pie de la hoja Resultado. Fuente y Nota quedan inmediatamente debajo de
    # la última fila de resultados, conservando texto y formato de la plantilla.
    # Las notas finales/exclusiones se agregan después con el mismo formato que
    # la fila Fuente y siempre alineadas a la izquierda.
    fila_fuente = ultima_fila_datos + 1
    fila_nota = fila_fuente + 1
    fila_final = fila_nota + 1

    for fila, texto_pie, estilo, altura in (
        (fila_fuente, pie["fuente"], pie["estilo_fuente"], pie["altura_fuente"]),
        (fila_nota, pie["nota"], pie["estilo_nota"], pie["altura_nota"]),
    ):
        resultado.merge_cells(start_row=fila, start_column=1,
                               end_row=fila, end_column=len(encabezados))
        celda = resultado.cell(fila, 1, texto_pie)
        _aplicar_estilo_celda(celda, estilo, alineacion="left")
        if altura:
            resultado.row_dimensions[fila].height = altura

    resultado.merge_cells(start_row=fila_final, start_column=1,
                           end_row=fila_final, end_column=len(encabezados))
    celda_final = resultado.cell(fila_final, 1, pie["nota_final"])
    # El usuario pidió que las notas finales adopten el formato de Fuente.
    _aplicar_estilo_celda(celda_final, pie["estilo_fuente"], alineacion="left")

    exclusiones = _exclusiones_calculo(interp)
    if not exclusiones:
        exclusiones = [
            "Categorías/códigos excluidos del cálculo: no se identifican "
            "exclusiones explícitas en el diccionario para las variables utilizadas."
        ]
    for offset, linea in enumerate(exclusiones, start=1):
        fila = fila_final + offset
        resultado.merge_cells(start_row=fila, start_column=1,
                               end_row=fila, end_column=len(encabezados))
        celda = resultado.cell(fila, 1, linea)
        _aplicar_estilo_celda(celda, pie["estilo_fuente"], alineacion="left")
        resultado.row_dimensions[fila].height = 32

    _agregar_hoja_piramides(wb, respuesta)
    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida, nombre_archivo(titulo)
