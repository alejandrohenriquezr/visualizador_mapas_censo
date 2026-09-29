# -*- coding: utf-8 -*-
"""Regresiones V29: selección explícita de denominadores y metadatos Excel."""
from __future__ import annotations

import json
import sys
import tempfile
import types
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

# El entorno de construcción no necesita DuckDB para inspeccionar SQL.
sys.modules.setdefault("duckdb", types.SimpleNamespace(connect=lambda *a, **k: None))

import data_engine
from denominator_resolver import AmbiguedadDenominador
from excel_export import crear_excel
from nl_parser import interpretar_consulta
from presentation import descripcion_cotidiana


Q_SIN_INTERNET = (
    "porcentaje de personas mayores en viviendas particulares sin acceso a internet"
)
Q_RELATIVA = (
    "Distribución relativa del total de personas mayores en viviendas particulares "
    "sin acceso a internet en el hogar, según sexo y edad quinquenal"
)


def _hojas(nodo):
    if not isinstance(nodo, dict):
        return []
    if nodo.get("op") in {"and", "or"}:
        out = []
        for item in nodo.get("args") or []:
            out.extend(_hojas(item))
        return out
    if nodo.get("op") == "not":
        return _hojas(nodo.get("arg"))
    return [nodo]


class DenominadoresParserV29Tests(unittest.TestCase):
    def test_porcentaje_sin_denominador_pregunta(self):
        with self.assertRaises(AmbiguedadDenominador) as ctx:
            interpretar_consulta(Q_SIN_INTERNET)
        opciones = ctx.exception.opciones
        self.assertEqual(
            [x["id"] for x in opciones],
            ["total_personas_mayores", "total_personas"],
        )
        self.assertIn("denominador", str(ctx.exception).lower())

    def test_seleccion_total_personas_mayores_se_registra(self):
        i = interpretar_consulta(
            Q_SIN_INTERNET,
            denominador_seleccionado="total_personas_mayores",
        )
        self.assertEqual(i["denominador_id"], "total_personas_mayores")
        self.assertEqual(i["denominador_confirmacion"], "seleccion_usuario")
        self.assertTrue(i["denominador_requirio_confirmacion"])
        self.assertIn("60 años o más", i["universo"])
        self.assertIn("total de personas de 60 años o más", i["formula"])
        pct = i["plan_cruce"]["porcentaje"]
        self.assertTrue(pct["denominador_personalizado"])
        hojas = _hojas(pct["denominador_ast"])
        self.assertEqual(len(hojas), 1)
        self.assertEqual(hojas[0]["variable"], "edad")
        self.assertEqual(hojas[0]["minimo"], 60)

    def test_seleccion_total_personas_elimina_filtros_del_denominador(self):
        i = interpretar_consulta(
            Q_SIN_INTERNET,
            denominador_seleccionado="total_personas",
        )
        self.assertEqual(i["denominador_id"], "total_personas")
        self.assertIsNone(i["plan_cruce"]["porcentaje"]["denominador_ast"])
        self.assertEqual(i["universo"], "Total de personas del territorio.")

    def test_denominador_explicito_personas_mayores_se_respeta(self):
        i = interpretar_consulta(
            Q_SIN_INTERNET + " sobre el total de personas mayores"
        )
        self.assertEqual(i["denominador_id"], "total_personas_mayores")
        self.assertEqual(i["denominador_confirmacion"], "explicito_en_pregunta")
        self.assertFalse(i["denominador_requirio_confirmacion"])

    def test_denominador_explicito_total_personas_se_respeta(self):
        i = interpretar_consulta(Q_SIN_INTERNET + " sobre el total de personas")
        self.assertEqual(i["denominador_id"], "total_personas")
        self.assertEqual(i["denominador_confirmacion"], "explicito_en_pregunta")

    def test_distribucion_relativa_del_total_ya_contiene_denominador(self):
        i = interpretar_consulta(Q_RELATIVA)
        self.assertTrue(i["denominador_definido"])
        self.assertEqual(i["operacion"], "porcentaje")
        self.assertIn("total de personas de 60 años o más", i["formula"])

    def test_tasa_metodologicamente_definida_no_pregunta(self):
        i = interpretar_consulta("tasa de desocupación región Metropolitana")
        self.assertEqual(i["operacion"], "razon")
        self.assertIn("desocupadas", i["formula"])
        self.assertNotIn("denominador_candidatos", i)

    def test_porcentaje_simple_con_denominador_explicito_total_entidad(self):
        i = interpretar_consulta("porcentaje de mujeres sobre el total de personas por región")
        self.assertEqual(i["denominador_id"], "total_entidad")
        self.assertTrue(i["denominador_entidad_total"])
        self.assertEqual(i["denominador_confirmacion"], "explicito_en_pregunta")

    def test_porcentaje_compuesto_generico_pregunta_base(self):
        with self.assertRaises(AmbiguedadDenominador) as ctx:
            interpretar_consulta("porcentaje de mujeres con discapacidad por región")
        ids = {x["id"] for x in ctx.exception.opciones}
        self.assertIn("subgrupo_contexto", ids)
        self.assertIn("total_entidad", ids)

    def test_seleccion_invalida_se_rechaza(self):
        with self.assertRaisesRegex(ValueError, "denominador seleccionado"):
            interpretar_consulta(Q_SIN_INTERNET, denominador_seleccionado="no_existe")


class DenominadoresSQLV29Tests(unittest.TestCase):
    def setUp(self):
        self.original_ruta = data_engine._ruta_parquet
        data_engine._ruta_parquet = lambda tabla: Path(f"C:/tmp/{tabla}.parquet")

    def tearDown(self):
        data_engine._ruta_parquet = self.original_ruta

    def test_numerador_conserva_condiciones_y_denominador_mayores_solo_edad(self):
        i = interpretar_consulta(
            Q_SIN_INTERNET,
            denominador_seleccionado="total_personas_mayores",
        )
        sql_num, _ = data_engine._construir_sql_cruce(i)
        sql_den, _ = data_engine._construir_sql_denominador_cruce(i)
        n = " ".join(sql_num.split())
        d = " ".join(sql_den.split())
        self.assertIn('TRY_CAST(p."edad" AS DOUBLE) >= 60', n)
        self.assertIn('CAST(p."tipo_operativo" AS VARCHAR) IN (\'2\')', n)
        self.assertIn('p15d_serv_internet_fija', n)
        self.assertIn('TRY_CAST(p."edad" AS DOUBLE) >= 60', d)
        self.assertNotIn('tipo_operativo', d)
        self.assertNotIn('p15d_serv_internet_fija', d)
        self.assertIn('COUNT(DISTINCT', d)

    def test_denominador_total_personas_no_tiene_where_tematico(self):
        i = interpretar_consulta(
            Q_SIN_INTERNET,
            denominador_seleccionado="total_personas",
        )
        sql_den, _ = data_engine._construir_sql_denominador_cruce(i)
        d = " ".join(sql_den.split())
        self.assertNotIn("WHERE", d)
        self.assertNotIn("JOIN read_parquet", d)

    def test_simple_total_entidad_usa_true_como_base(self):
        i = interpretar_consulta("porcentaje de mujeres sobre el total de personas por región")
        sql = data_engine._construir_sql(
            i, data_engine.NIVELES["region"], Path("C:/tmp/personas.parquet")
        )
        self.assertIn("SUM(CASE WHEN TRUE THEN 1 ELSE 0 END)", sql)

    def test_jefatura_cambia_denominador_segun_eleccion(self):
        q = "porcentaje de personas mayores de 60 años jefas de hogar"
        mayores = interpretar_consulta(q, denominador_seleccionado="total_personas_mayores")
        todas = interpretar_consulta(q, denominador_seleccionado="total_personas")
        sql_m = " ".join(data_engine._construir_sql_indicador_censal_v27(
            mayores, data_engine.NIVELES[mayores["nivel_geografico"]], Path("C:/tmp/personas.parquet")
        ).split())
        sql_t = " ".join(data_engine._construir_sql_indicador_censal_v27(
            todas, data_engine.NIVELES[todas["nivel_geografico"]], Path("C:/tmp/personas.parquet")
        ).split())
        self.assertIn('SUM(CASE WHEN TRY_CAST("edad" AS DOUBLE) > 60 AND TRY_CAST("edad" AS DOUBLE) <= 120 THEN 1 ELSE 0 END)', sql_m)
        self.assertIn("COUNT(*)", sql_t)


class ExcelV29Tests(unittest.TestCase):
    @staticmethod
    def _plantilla(d):
        p = Path(d) / "plantilla.xlsx"
        wb = Workbook()
        nota = wb.active
        nota.title = "Nota"
        nota["B5"] = "Fecha"
        resultado = wb.create_sheet("Resultado")
        resultado["A1"] = "Título"
        resultado["A22"] = "Nota final"
        wb.save(p)
        return p

    def test_universo_formula_exactamente_en_b_e(self):
        with tempfile.TemporaryDirectory() as d:
            p = self._plantilla(d)
            i = interpretar_consulta(
                Q_SIN_INTERNET,
                denominador_seleccionado="total_personas_mayores",
            )
            respuesta = {
                "tipo_visualizacion": "mapa",
                "nivel_geografico": "region",
                "interpretacion": {
                    "descripcion": descripcion_cotidiana(i),
                    "tabla": i["tabla"],
                    "variable": i["variable"],
                    "operacion": i["operacion"],
                    "nivel_geografico": "region",
                    "entidad_objetivo": "personas",
                    "universo": i["universo"],
                    "formula": i["formula"],
                },
                "datos": [{"codigo": 13, "nombre": "Metropolitana de Santiago", "valor": 11.2}],
            }
            contenido, _ = crear_excel(p, respuesta, Q_SIN_INTERNET, "SELECT 1")
            wb = load_workbook(contenido)
            nota = wb["Nota"]
            self.assertIsNone(nota["A20"].value)
            self.assertIsNone(nota["A21"].value)
            self.assertEqual(nota["B20"].value, "Universo")
            self.assertEqual(nota["E20"].value, i["universo"])
            self.assertEqual(nota["B21"].value, "Fórmula")
            self.assertEqual(nota["E21"].value, i["formula"])
            self.assertTrue(nota["B20"].font.bold)
            self.assertTrue(nota["E21"].font.bold)
            self.assertTrue(nota["E20"].alignment.wrap_text)
            self.assertTrue(nota["E21"].alignment.wrap_text)
            self.assertEqual(nota["E20"].alignment.vertical, "top")
            self.assertEqual(nota["E21"].alignment.vertical, "top")
            self.assertIsNone(nota["C20"].value)
            self.assertIsNone(nota["C21"].value)


class ContratoHTTPV29Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Importación tardía porque main carga data_engine.
        import main
        cls.main = main

    def test_request_acepta_denominador(self):
        req = self.main.ConsultaRequest(
            pregunta=Q_SIN_INTERNET,
            denominador_seleccionado="total_personas_mayores",
        )
        self.assertEqual(req.denominador_seleccionado, "total_personas_mayores")

    def test_endpoint_ambiguo_devuelve_409_y_opciones(self):
        resp = self.main.consulta(self.main.ConsultaRequest(pregunta=Q_SIN_INTERNET))
        self.assertEqual(resp.status_code, 409)
        data = json.loads(resp.body)
        self.assertEqual(data["tipo"], "seleccion_denominador")
        self.assertEqual(
            [x["id"] for x in data["opciones"]],
            ["total_personas_mayores", "total_personas"],
        )
        self.assertIn("seleccion_previa", data)


if __name__ == "__main__":
    unittest.main()
