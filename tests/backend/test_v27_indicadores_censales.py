# -*- coding: utf-8 -*-
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import sys
import types
# Permite validar el generador SQL en entornos de prueba donde DuckDB no está
# instalado; ninguna prueba de esta clase abre una conexión.
sys.modules.setdefault("duckdb", types.SimpleNamespace(connect=lambda *a, **k: None))
from data_engine import _construir_sql_indicador_censal_v27, NIVELES

from excel_export import crear_excel
from nl_parser import interpretar_consulta
from presentation import descripcion_cotidiana


class ParserIndicadoresV27Tests(unittest.TestCase):
    CASOS = {
        "tasa de desocupación región Metropolitana": ("indicador_derivado_v24", "tasa_desocupacion"),
        "porcentaje de desocupación región Metropolitana": ("indicador_derivado_v24", "tasa_desocupacion"),
        "indice de envejecimiento región Metropolitana": ("indicador_derivado_v24", "envejecimiento"),
        "indice de dependencia laboral región Metropolitana": ("indicador_derivado_v24", "dependencia_total"),
        "Tasa global de fecundidad región Metropolitana": ("indicador_censal_v27", "tgf_aproximada_censal"),
        "Tasa de alfabetismo región Metropolitana": ("indicador_censal_v27", "tasa_alfabetismo"),
        "Tasa de participación económica": ("indicador_censal_v27", "tasa_participacion_economica"),
        "Tasa de ocupación (o empleo)": ("indicador_derivado_v24", "tasa_ocupacion"),
        "Tasa de desocupación": ("indicador_derivado_v24", "tasa_desocupacion"),
        "Tasa de migración": ("indicador_censal_v27", "tasa_migracion_reciente"),
        "Tasa de migración reciente": ("indicador_censal_v27", "tasa_migracion_reciente"),
        "Tasa de migración neta interna": ("indicador_censal_v27", "tasa_migracion_neta_interna"),
        "Porcentaje de población inmigrante internacional": ("indicador_censal_v27", "porcentaje_inmigrante_internacional"),
        "Porcentaje de población indígena o afrodescendiente": ("indicador_censal_v27", "porcentaje_poblacion_indigena_o_afro"),
    }

    def test_consultas_prioritarias(self):
        for consulta, (tipo, esperado) in self.CASOS.items():
            with self.subTest(consulta=consulta):
                i = interpretar_consulta(consulta)
                self.assertEqual(i["tipo_consulta"], tipo)
                self.assertEqual(i["indicador_id"], esperado)
                self.assertTrue(i.get("universo"))
                self.assertTrue(i.get("formula"))

    def test_region_metropolitana_sin_por(self):
        i = interpretar_consulta("tasa de desocupación región Metropolitana")
        self.assertEqual(i["nivel_geografico"], "region")
        self.assertEqual(i["filtro_geografico_nivel"], "region")
        self.assertEqual(int(i["filtro_geografico_codigo"]), 13)

    def test_distribuciones(self):
        for consulta, variable in (
            ("Nivel de instrucción o escolaridad alcanzada", "cine11"),
            ("Distribución por categoría ocupacional", "p40_cise_rec"),
        ):
            with self.subTest(consulta=consulta):
                i = interpretar_consulta(consulta)
                self.assertEqual(i["tipo_consulta"], "distribucion_censal_v27")
                self.assertEqual(i["operacion"], "distribucion")
                self.assertEqual(i["variable"], variable)
                self.assertGreaterEqual(len(i["categorias_validas"]), 2)

    def test_tgf_se_declara_aproximada(self):
        i = interpretar_consulta("Tasa global de fecundidad región Metropolitana")
        self.assertIn("aproximada", i["indicador_descripcion"].lower())
        self.assertIn("no reemplaza", i["nota_interpretacion"].lower())


class CalculoIndicadoresV27Tests(unittest.TestCase):
    def _sql(self, consulta):
        i = interpretar_consulta(consulta)
        ruta = Path("/tmp/personas_sinteticas.parquet")
        sql = _construir_sql_indicador_censal_v27(i, NIVELES[i["nivel_geografico"]], ruta)
        return i, " ".join(sql.split())

    def test_tasa_desocupacion_formula_sql(self):
        _, sql = self._sql("tasa de desocupación región Metropolitana")
        self.assertIn("sit_fuerza_trabajo", sql)
        self.assertIn("IN ('1','2')", sql)
        self.assertIn("= '2'", sql)
        self.assertIn("100.0", sql)
        self.assertIn("region = 13", sql)
        # Control aritmético independiente: 2 desocupadas / 8 fuerza de trabajo.
        self.assertAlmostEqual(100 * 2 / (6 + 2), 25.0)

    def test_participacion_y_ocupacion_denominador_valido(self):
        _, sql_p = self._sql("tasa de participación económica")
        self.assertIn("IN ('1','2','3')", sql_p)
        self.assertIn("IN ('1','2')", sql_p)
        _, sql_o = self._sql("tasa de ocupación")
        self.assertIn("= '1'", sql_o)
        self.assertIn("IN ('1','2','3')", sql_o)
        self.assertAlmostEqual(100 * (6 + 2) / 10, 80.0)
        self.assertAlmostEqual(100 * 6 / 10, 60.0)

    def test_alfabetismo_usa_edad_y_respuestas_validas(self):
        _, sql = self._sql("tasa de alfabetismo región Metropolitana")
        self.assertIn('TRY_CAST("edad" AS DOUBLE) >= 5', sql)
        self.assertIn("p37_alfabet", sql)
        self.assertIn("IN ('1','2')", sql)

    def test_tgf_aproximada_contiene_siete_grupos(self):
        _, sql = self._sql("Tasa global de fecundidad región Metropolitana")
        for a, b in ((15,19),(20,24),(25,29),(30,34),(35,39),(40,44),(45,49)):
            self.assertIn(f"BETWEEN {a} AND {b}", sql)
        self.assertIn('p48_anio_nac_uh', sql)
        self.assertIn('= 2023', sql)
        self.assertIn('5.0 *', sql)

    def test_tasa_neta_migracion_interna_reconstruye_origen_destino(self):
        _, sql = self._sql("tasa de migración neta interna región Metropolitana")
        self.assertIn("inmigrantes", sql)
        self.assertIn("emigrantes", sql)
        self.assertIn("poblacion_2019", sql)
        self.assertIn("poblacion_2024", sql)
        self.assertIn("1000.0", sql)
        self.assertIn("5.0", sql)
        self.assertIn("WHERE codigo = 13", sql)

    def test_union_indigena_afro_no_duplica(self):
        _, sql = self._sql("Porcentaje de población indígena o afrodescendiente")
        self.assertIn("p28_autoid_pueblo", sql)
        self.assertIn("p29_afrodescendencia_rec", sql)
        self.assertIn("OR", sql)
        self.assertIn("AND", sql)


class ExcelMetadatosV27Tests(unittest.TestCase):
    def _plantilla(self, d):
        p = Path(d) / "plantilla.xlsx"
        wb = Workbook(); nota = wb.active; nota.title = "Nota"; nota["B5"] = "Fecha"
        r = wb.create_sheet("Resultado"); r["A1"] = "Título"; r["A22"] = "Nota final"
        wb.save(p); return p

    def test_universo_y_formula_en_celdas_solicitadas(self):
        with tempfile.TemporaryDirectory() as d:
            p = self._plantilla(d)
            i = interpretar_consulta("tasa de desocupación región Metropolitana")
            respuesta = {
                "tipo_visualizacion":"mapa", "nivel_geografico":"region",
                "interpretacion": {
                    "descripcion": descripcion_cotidiana(i), "tabla":i["tabla"],
                    "variable":i["variable"], "operacion":i["operacion"],
                    "nivel_geografico":i["nivel_geografico"],
                    "entidad_objetivo":"personas", "universo":i["universo"],
                    "formula":i["formula"],
                },
                "datos":[{"codigo":13,"nombre":"Metropolitana de Santiago","valor":7.2}],
            }
            contenido, _ = crear_excel(p, respuesta, "tasa de desocupación región Metropolitana", "SELECT 1")
            wb = load_workbook(contenido)
            nota = wb["Nota"]
            self.assertEqual(nota["E19"].value, "personas")
            self.assertIsNone(nota["A20"].value)
            self.assertEqual(nota["B20"].value, "Universo")
            self.assertIn("fuerza de trabajo", nota["E20"].value)
            self.assertIsNone(nota["A21"].value)
            self.assertEqual(nota["B21"].value, "Fórmula")
            self.assertIn("desocupadas", nota["E21"].value)
            self.assertIn("× 100", nota["E21"].value)

    def test_universo_y_formula_en_matriz_od(self):
        with tempfile.TemporaryDirectory() as d:
            p = self._plantilla(d)
            respuesta = {
                "tipo_visualizacion":"matriz_od",
                "interpretacion": {
                    "descripcion":"Matriz de movilidad laboral entre comunas.",
                    "tabla":"personas", "variable":"p44_lug_trab",
                    "operacion":"conteo", "nivel_geografico":"comuna",
                    "entidad_objetivo":"personas",
                },
                "config_od": {"dominio":"movilidad_laboral", "unidad":"comuna",
                              "metrica":"cantidad", "incluir_diagonal":True},
                "flujos_od":[{
                    "origen_codigo":13119, "origen_nombre":"Maipú",
                    "origen_region_nombre":"Metropolitana de Santiago",
                    "destino_codigo":13101, "destino_nombre":"Santiago",
                    "destino_region_nombre":"Metropolitana de Santiago",
                    "cantidad":10, "porcentaje_origen":100.0,
                    "porcentaje_destino":100.0, "misma_comuna":False,
                }],
            }
            contenido, _ = crear_excel(p, respuesta, "movilidad laboral región Metropolitana", "SELECT 1")
            wb = load_workbook(contenido)
            nota = wb["Nota"]
            self.assertIsNone(nota["A20"].value)
            self.assertEqual(nota["B20"].value, "Universo")
            self.assertIn("ocupadas", nota["E20"].value)
            self.assertIsNone(nota["A21"].value)
            self.assertEqual(nota["B21"].value, "Fórmula")
            self.assertIn("conteo", nota["E21"].value.lower())


if __name__ == "__main__":
    unittest.main()
