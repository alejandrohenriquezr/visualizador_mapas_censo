# -*- coding: utf-8 -*-
"""Regresiones reportadas tras v25: OD, Mapuzugún y exportación Excel."""

from __future__ import annotations

import sys
import tempfile
import types
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from nl_parser import interpretar_consulta
from excel_export import crear_excel


class MovilidadLaboralV26Tests(unittest.TestCase):
    def _assert_region_matriz(self, consulta, codigo):
        i = interpretar_consulta(consulta)
        self.assertEqual(i["tipo_consulta"], "movilidad_laboral_fase3")
        self.assertEqual(i["movilidad_indicador"], "matriz_od_comunal")
        self.assertEqual(i["nivel_geografico"], "comuna")
        self.assertEqual(i["movilidad_region_origen_codigo"], codigo)
        self.assertEqual(i["movilidad_region_destino_codigo"], codigo)
        return i

    def test_por_comunas_region_metropolitana_es_matriz_od(self):
        self._assert_region_matriz(
            "Movilidad laboral por comunas de la región Metropolitana", 13
        )

    def test_movilidad_region_metropolitana_es_matriz_od(self):
        self._assert_region_matriz("Movilidad laboral región Metropolitana", 13)

    def test_region_abreviada_aysen(self):
        self._assert_region_matriz("Movilidad laboral región de Aysén", 11)

    def test_matriz_historica_de_las_comunas_conserva_solo_origen(self):
        i = interpretar_consulta("matriz de movilidad laboral de las comunas de La Araucanía")
        self.assertEqual(i["movilidad_region_origen_codigo"], 9)
        self.assertIsNone(i.get("movilidad_region_destino_codigo"))


# data_engine importa duckdb en carga; en el entorno de construcción puede no
# estar instalado. El stub basta para inspeccionar SQL sin ejecutar consultas.
try:
    import duckdb  # noqa: F401
except ModuleNotFoundError:
    sys.modules["duckdb"] = types.SimpleNamespace(connect=lambda *a, **k: None)

import data_engine  # noqa: E402


class MapuzugunV26Tests(unittest.TestCase):
    def test_sql_porcentaje_admite_codigo_numerico_decimal(self):
        i = interpretar_consulta(
            "porcentaje de personas que hablan o entienden Mapuzugún por región"
        )
        sql = data_engine._construir_sql(
            i, data_engine.NIVELES["region"], Path("C:/tmp/personas.parquet")
        )
        self.assertIn('TRY_CAST("p30_lengua_indigena" AS INTEGER) IN (1)', sql)
        self.assertIn(
            'TRY_CAST("p30_lengua_indigena" AS INTEGER) IN (1, 2, 3, 4, 5, 6, 7, 8, 9)',
            sql,
        )
        self.assertNotIn("Distribución", i.get("indicador_descripcion", ""))

    def test_sql_conteo_mapuzugun_usa_codigo_numerico(self):
        i = interpretar_consulta("personas que hablan o entienden Mapuzugún por región")
        sql = data_engine._construir_sql(
            i, data_engine.NIVELES["region"], Path("C:/tmp/personas.parquet")
        )
        self.assertIn(
            'SUM(CASE WHEN TRY_CAST("p30_lengua_indigena" AS INTEGER) = 1 THEN 1 ELSE 0 END)',
            sql,
        )


class ExcelEntidadObjetivoV26Tests(unittest.TestCase):
    def _plantilla(self, directorio: Path) -> Path:
        ruta = directorio / "plantilla.xlsx"
        wb = Workbook()
        nota = wb.active
        nota.title = "Nota"
        nota["B5"] = "Fecha"
        resultado = wb.create_sheet("Resultado")
        resultado["A1"] = "Título"
        resultado["A22"] = "Nota final"
        wb.save(ruta)
        return ruta

    def test_excel_mapa_escribe_entidad_objetivo_en_e19(self):
        with tempfile.TemporaryDirectory() as tmp:
            plantilla = self._plantilla(Path(tmp))
            respuesta = {
                "tipo_visualizacion": "mapa",
                "nivel_geografico": "region",
                "interpretacion": {
                    "descripcion": "Cantidad de personas según regiones.",
                    "tabla": "personas",
                    "variable": "edad",
                    "operacion": "conteo",
                    "nivel_geografico": "region",
                },
                "datos": [{"codigo": 13, "nombre": "Metropolitana de Santiago", "valor": 10}],
            }
            contenido, _ = crear_excel(plantilla, respuesta, "personas por región", "SELECT 1")
            wb = load_workbook(contenido)
            self.assertEqual(wb["Nota"]["B19"].value, "Entidad objetivo")
            self.assertEqual(wb["Nota"]["E19"].value, "personas")

    def test_excel_od_escribe_entidad_objetivo_en_e19(self):
        with tempfile.TemporaryDirectory() as tmp:
            plantilla = self._plantilla(Path(tmp))
            respuesta = {
                "tipo_visualizacion": "matriz_od",
                "interpretacion": {
                    "descripcion": "Matriz de movilidad laboral entre comunas.",
                    "tabla": "personas",
                    "variable": "p44_lug_trab",
                    "operacion": "conteo",
                    "nivel_geografico": "comuna",
                },
                "config_od": {
                    "dominio": "movilidad_laboral",
                    "unidad": "comuna",
                    "metrica": "cantidad",
                    "incluir_diagonal": True,
                },
                "flujos_od": [{
                    "origen_codigo": 13119,
                    "origen_nombre": "Maipú",
                    "origen_region_nombre": "Metropolitana de Santiago",
                    "destino_codigo": 13101,
                    "destino_nombre": "Santiago",
                    "destino_region_nombre": "Metropolitana de Santiago",
                    "cantidad": 100,
                    "porcentaje_origen": 100.0,
                    "porcentaje_destino": 100.0,
                    "misma_comuna": False,
                }],
            }
            contenido, _ = crear_excel(plantilla, respuesta, "movilidad laboral región Metropolitana", "SELECT 1")
            wb = load_workbook(contenido)
            self.assertEqual(wb["Nota"]["B19"].value, "Entidad objetivo")
            self.assertEqual(wb["Nota"]["E19"].value, "personas")


if __name__ == "__main__":
    unittest.main()
