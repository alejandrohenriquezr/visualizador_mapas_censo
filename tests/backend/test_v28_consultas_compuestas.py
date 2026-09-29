# -*- coding: utf-8 -*-
"""Regresiones V28 para consultas compuestas y metadatos Excel."""
import sys
import tempfile
import types
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

# El generador SQL se prueba sin abrir una conexión DuckDB.
sys.modules.setdefault("duckdb", types.SimpleNamespace(connect=lambda *a, **k: None))

import data_engine
from excel_export import crear_excel
from nl_parser import interpretar_consulta
from denominator_resolver import AmbiguedadDenominador
from presentation import descripcion_cotidiana


Q_TENENCIA_PERSONAS = (
    "personas que viven en viviendas arrendadas sin contrato, ocupadas de hecho "
    "o pertenecientes a una propiedad en sucesión o litigio"
)
Q_TENENCIA_HOGARES = (
    "Hogares con viviendas arrendadas sin contrato, ocupadas de hecho o "
    "pertenecientes a una propiedad en sucesión o litigio"
)
Q_MAYORES_REL = (
    "Distribución relativa del total de personas mayores en viviendas particulares "
    "sin acceso a internet en el hogar, según sexo y edad quinquenal"
)
Q_MAYORES = (
    "Distribución de personas mayores en viviendas particulares sin acceso a internet "
    "en el hogar, según sexo y edad quinquenal"
)


def _hojas(nodo):
    if not isinstance(nodo, dict):
        return []
    if nodo.get("op") in {"and", "or"}:
        out = []
        for item in nodo.get("args") or []:
            out.extend(_hojas(item))
        return out
    if nodo.get("op") == "not":
        return _hojas(nodo.get("arg"))
    return [nodo]


class ParserConsultasCompuestasV28Tests(unittest.TestCase):
    def test_tenencia_personas_no_se_confunde_con_ocupacion_laboral(self):
        i = interpretar_consulta(Q_TENENCIA_PERSONAS)
        self.assertEqual(i["tipo_consulta"], "cruce")
        self.assertEqual(i["plan_cruce"]["entidad_objetivo"], "personas")
        filtros = _hojas(i["plan_cruce"]["filtro_ast"])
        ten = next(f for f in filtros if f.get("variable") == "p12_tenencia_viv")
        self.assertEqual(ten["valores"], ["4", "8", "9"])
        self.assertFalse(any(f.get("variable") == "sit_fuerza_trabajo" for f in filtros))
        self.assertIn("arrendadas sin contrato", descripcion_cotidiana(i).lower())

    def test_tenencia_hogares_cuenta_hogares(self):
        i = interpretar_consulta(Q_TENENCIA_HOGARES)
        self.assertEqual(i["tipo_consulta"], "cruce")
        self.assertEqual(i["plan_cruce"]["entidad_objetivo"], "hogares")
        filtros = _hojas(i["plan_cruce"]["filtro_ast"])
        ten = next(f for f in filtros if f.get("variable") == "p12_tenencia_viv")
        self.assertEqual(ten["valores"], ["4", "8", "9"])

    def test_ocupados_arrendados_sigue_siendo_consulta_laboral(self):
        i = interpretar_consulta("ocupados que viven en viviendas arrendadas por comuna")
        plan = i["plan_cruce"]
        filtros = plan.get("filtros") or _hojas(plan.get("filtro_ast"))
        self.assertTrue(any(f.get("variable") == "sit_fuerza_trabajo" and f.get("valores") == ["1"] for f in filtros))
        self.assertTrue(any(f.get("variable") == "p12_tenencia_viv" and f.get("valores") == ["3", "4"] for f in filtros))

    def test_personas_mayores_equivale_a_60_o_mas_en_consulta_compuesta(self):
        i = interpretar_consulta(Q_MAYORES_REL)
        self.assertEqual(i["tipo_consulta"], "cruce")
        self.assertEqual(i["operacion"], "porcentaje")
        p = i["plan_cruce"]
        self.assertEqual(p["entidad_objetivo"], "personas")
        self.assertEqual([(d["tabla"], d["variable"]) for d in p["dimensiones"]],
                         [("personas", "sexo"), ("personas", "edad_quinquenal")])
        self.assertEqual(p["porcentaje"]["base"], "total")
        filtros = _hojas(p["filtro_ast"])
        edad = next(f for f in filtros if f.get("variable") == "edad")
        self.assertEqual(edad["minimo"], 60)
        self.assertTrue(edad["incluir_minimo"])
        self.assertTrue(any(f.get("variable") == "tipo_operativo" and f.get("valores") == ["2"] for f in filtros))
        for var in ("p15d_serv_internet_fija", "p15e_serv_internet_movil", "p15f_serv_internet_satelital"):
            self.assertTrue(any(f.get("variable") == var and f.get("valores") == ["2"] for f in filtros))
        self.assertIn("60 años o más", i["universo"])

    def test_distribucion_no_relativa_conserva_dos_dimensiones(self):
        i = interpretar_consulta(Q_MAYORES)
        self.assertEqual(i["operacion"], "distribucion")
        self.assertNotIn("porcentaje", i["plan_cruce"])
        self.assertEqual(len(i["plan_cruce"]["dimensiones"]), 2)

    def test_composicion_hogar_resuelve_tipologia(self):
        for q in (
            "Distribución relativa del total de hogares según composición del hogar",
            "Distribución de hogares según composición del hogar",
        ):
            with self.subTest(q=q):
                i = interpretar_consulta(q)
                self.assertEqual(i["tabla"], "hogares")
                self.assertEqual(i["variable"], "tipologia_hogar")
                self.assertEqual(i["operacion"], "distribucion")
                self.assertIn("tipología de hogar válida", i["universo"])

    def test_porcentaje_personas_mayores_jefatura_pregunta_denominador(self):
        with self.assertRaises(AmbiguedadDenominador) as ctx:
            interpretar_consulta("porcentaje de personas mayores de 60 años jefas de hogar")
        ids = [x["id"] for x in ctx.exception.opciones]
        self.assertEqual(ids, ["total_personas_mayores", "total_personas"])

    def test_porcentaje_personas_mayores_jefatura_con_denominador(self):
        i = interpretar_consulta(
            "porcentaje de personas mayores de 60 años jefas de hogar",
            denominador_seleccionado="total_personas_mayores",
        )
        self.assertEqual(i["tipo_consulta"], "indicador_censal_v27")
        self.assertEqual(i["indicador_id"], "porcentaje_personas_mayores_jefatura")
        self.assertEqual(i["edad_umbral"], 60)
        self.assertFalse(i["edad_inclusiva"])
        self.assertEqual(i["denominador_id"], "total_personas_mayores")
        self.assertIn("mayores de 60 años", i["formula"])
        self.assertIn("jefas o jefes", descripcion_cotidiana(i).lower())

    def test_personas_mayores_sin_numero_se_define_60_o_mas(self):
        i = interpretar_consulta(
            "porcentaje de personas mayores jefas de hogar",
            denominador_seleccionado="total_personas_mayores",
        )
        self.assertEqual(i["edad_umbral"], 60)
        self.assertTrue(i["edad_inclusiva"])


class SQLConsultasCompuestasV28Tests(unittest.TestCase):
    def setUp(self):
        self.original_ruta = data_engine._ruta_parquet
        data_engine._ruta_parquet = lambda tabla: Path(f"C:/tmp/{tabla}.parquet")

    def tearDown(self):
        data_engine._ruta_parquet = self.original_ruta

    def test_sql_tenencia_personas_usa_join_y_codigos_exactos(self):
        i = interpretar_consulta(Q_TENENCIA_PERSONAS)
        sql, _ = data_engine._construir_sql_cruce(i)
        sql = " ".join(sql.split())
        self.assertIn("JOIN read_parquet('C:/tmp/hogares.parquet') h", sql)
        self.assertIn('CAST(h."p12_tenencia_viv" AS VARCHAR) IN (\'4\', \'8\', \'9\')', sql)
        self.assertNotIn("sit_fuerza_trabajo", sql)
        self.assertIn("COUNT(DISTINCT", sql)

    def test_sql_mayores_sin_internet_aplica_todo_el_universo(self):
        i = interpretar_consulta(Q_MAYORES_REL)
        sql, _ = data_engine._construir_sql_cruce(i)
        sql = " ".join(sql.split())
        self.assertIn('TRY_CAST(p."edad" AS DOUBLE) >= 60', sql)
        self.assertIn('CAST(p."tipo_operativo" AS VARCHAR) IN (\'2\')', sql)
        self.assertIn('CAST(h."p15d_serv_internet_fija" AS VARCHAR) IN (\'2\')', sql)
        self.assertIn('CAST(h."p15e_serv_internet_movil" AS VARCHAR) IN (\'2\')', sql)
        self.assertIn('CAST(h."p15f_serv_internet_satelital" AS VARCHAR) IN (\'2\')', sql)
        self.assertIn('CAST(p."sexo" AS VARCHAR) AS d0', sql)
        self.assertIn('CAST(p."edad_quinquenal" AS VARCHAR) AS d1', sql)

    def test_sql_porcentaje_jefatura_mayores_de_60(self):
        i = interpretar_consulta(
            "porcentaje de personas mayores de 60 años jefas de hogar",
            denominador_seleccionado="total_personas_mayores",
        )
        sql = data_engine._construir_sql_indicador_censal_v27(
            i, data_engine.NIVELES[i["nivel_geografico"]], Path("C:/tmp/personas.parquet")
        )
        sql = " ".join(sql.split())
        self.assertIn('TRY_CAST("edad" AS DOUBLE) > 60', sql)
        self.assertIn('CAST("tipo_operativo" AS VARCHAR) = \'2\'', sql)
        self.assertIn('CAST("parentesco" AS VARCHAR) = \'1\'', sql)
        self.assertIn("100.0", sql)


class ExcelNotaV28Tests(unittest.TestCase):
    @staticmethod
    def _plantilla(d):
        p = Path(d) / "plantilla.xlsx"
        wb = Workbook()
        nota = wb.active
        nota.title = "Nota"
        nota["B5"] = "Fecha"
        resultado = wb.create_sheet("Resultado")
        resultado["A1"] = "Título"
        resultado["A22"] = "Nota final"
        wb.save(p)
        return p

    def test_universo_y_formula_se_escriben_en_columnas_b_e(self):
        with tempfile.TemporaryDirectory() as d:
            p = self._plantilla(d)
            i = interpretar_consulta(
                "porcentaje de personas mayores de 60 años jefas de hogar",
                denominador_seleccionado="total_personas_mayores",
            )
            respuesta = {
                "tipo_visualizacion": "mapa",
                "nivel_geografico": "region",
                "interpretacion": {
                    "descripcion": descripcion_cotidiana(i),
                    "tabla": i["tabla"], "variable": i["variable"],
                    "operacion": i["operacion"], "nivel_geografico": "region",
                    "entidad_objetivo": "personas",
                    "universo": i["universo"], "formula": i["formula"],
                },
                "datos": [{"codigo": 13, "nombre": "Metropolitana de Santiago", "valor": 31.5}],
            }
            contenido, _ = crear_excel(p, respuesta, "consulta", "SELECT 1")
            wb = load_workbook(contenido)
            nota = wb["Nota"]
            self.assertEqual(nota["E19"].value, "personas")
            self.assertIsNone(nota["A20"].value)
            self.assertEqual(nota["B20"].value, "Universo")
            self.assertIn("mayores de 60", nota["E20"].value)
            self.assertIsNone(nota["A21"].value)
            self.assertEqual(nota["B21"].value, "Fórmula")
            self.assertIn("× 100", nota["E21"].value)


if __name__ == "__main__":
    unittest.main()
